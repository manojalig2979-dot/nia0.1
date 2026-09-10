import os
import sys
import subprocess
import webbrowser
from datetime import datetime
import mss
import mss.tools
from docx import Document

class DesktopTools:
    def __init__(self, projects_dir: str):
        self.projects_dir = projects_dir

    def open_application(self, app_name: str) -> str:
        """Opens installed applications like notepad, chrome, calculator, etc."""
        app_name_clean = app_name.lower().strip()
        system = sys.platform
        
        apps_map = {
            "notepad": "notepad.exe" if system == "win32" else "gedit",
            "chrome": "chrome" if system != "win32" else "start chrome",
            "calculator": "calc.exe" if system == "win32" else "gnome-calculator",
            "vs code": "code",
            "vscode": "code",
            "taskmgr": "start taskmgr",
            "explorer": "start explorer",
            "cmd": "start cmd",
            "mspaint": "start mspaint",
            "paint": "start mspaint",
            "spotify": "start spotify"
        }

        command = apps_map.get(app_name_clean, app_name_clean)
        try:
            if system == "win32":
                subprocess.Popen(command, shell=True)
            else:
                subprocess.Popen([command], shell=True)
            return f"Successfully opened {app_name}."
        except Exception as e:
            return f"Failed to open {app_name}: {str(e)}"

    def open_vscode_project(self, project_name: str) -> str:
        """Locates a project in the designated projects folder and opens it in VS Code."""
        target_path = os.path.join(self.projects_dir, project_name)
        if not os.path.exists(target_path):
            # Attempt case-insensitive or partial lookup
            matches = [d for d in os.listdir(self.projects_dir) if project_name.lower() in d.lower()]
            if matches:
                target_path = os.path.join(self.projects_dir, matches[0])
            else:
                # If directory doesn't exist, create it
                os.makedirs(target_path, exist_ok=True)
        
        try:
            subprocess.Popen(["code", target_path], shell=True)
            return f"Opened project '{project_name}' in VS Code at path: {target_path}"
        except Exception as e:
            return f"Could not launch VS Code: {str(e)}"

    def open_url(self, url: str) -> str:
        """Opens any website URL in the user's default browser."""
        if not url.startswith("http"):
            url = "https://" + url
        webbrowser.open(url)
        return f"Opened {url} in browser."

    def take_screenshot(self, filename_prefix="screenshot") -> str:
        """Captures all displays and saves the image with a timestamp."""
        output_dir = os.path.join(os.path.expanduser("~"), "Pictures", "NiaScreenshots")
        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filepath = os.path.join(output_dir, f"{filename_prefix}_{timestamp}.png")
        
        with mss.mss() as sct:
            sct.shot(mon=-1, output=filepath)
        return filepath

    def draft_document(self, title: str, doc_type: str, content: str) -> str:
        """Drafts structured letters, memos, or project specs in Word docx format."""
        docs_dir = os.path.join(os.path.expanduser("~"), "Documents", "NiaDrafts")
        os.makedirs(docs_dir, exist_ok=True)
        
        clean_title = "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).strip()
        filename = f"{clean_title or 'Draft'}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        filepath = os.path.join(docs_dir, filename)

        doc = Document()
        doc.add_heading(title, level=0)
        doc.add_paragraph(f"Document Type: {doc_type.upper()}")
        doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
        doc.add_paragraph("-" * 40)
        doc.add_paragraph(content)

        doc.save(filepath)
        return f"Document drafted and saved successfully at: {filepath}"

    def system_power_control(self, action: str) -> str:
        """Controls system shutdown or reboot."""
        action = action.lower().strip()
        if sys.platform != "win32":
            return "Shutdown commands are only configured for Windows currently."
        
        if "shutdown" in action:
            os.system("shutdown /s /t 10")
            return "System will shut down in 10 seconds. Save your work."
        elif "restart" in action or "reboot" in action:
            os.system("shutdown /r /t 10")
            return "System will restart in 10 seconds."
        elif "cancel" in action:
            os.system("shutdown /a")
            return "System shutdown/restart canceled."
        return "Unknown power command. Specify 'shutdown', 'restart', or 'cancel'."

    def play_in_windows_media_player(self, song_name: str = "") -> str:
        """Plays songs or audio in Windows Media Player (Default music player)."""
        wmplayer_path = os.path.expandvars(r"%ProgramFiles%\Windows Media Player\wmplayer.exe")
        if not os.path.exists(wmplayer_path):
            wmplayer_path = os.path.expandvars(r"%ProgramFiles(x86)%\Windows Media Player\wmplayer.exe")

        music_dirs = [
            os.path.join(os.path.expanduser("~"), "Music"),
            os.path.join(os.path.expanduser("~"), "Downloads"),
            self.projects_dir
        ]

        audio_exts = (".mp3", ".wav", ".m4a", ".flac", ".wma", ".aac", ".ogg")
        matched_file = None

        if song_name:
            clean_query = song_name.lower().strip()
            for mdir in music_dirs:
                if os.path.exists(mdir):
                    for root, _, files in os.walk(mdir):
                        for f in files:
                            if f.lower().endswith(audio_exts):
                                if any(word in f.lower() for word in clean_query.split() if len(word) > 2):
                                    matched_file = os.path.join(root, f)
                                    break
                        if matched_file:
                            break
                if matched_file:
                    break

        if matched_file and os.path.exists(matched_file):
            if os.path.exists(wmplayer_path):
                subprocess.Popen([wmplayer_path, matched_file])
            else:
                os.startfile(matched_file)
            return f"Windows Media Player me '{os.path.basename(matched_file)}' play kar diya hai."
        else:
            if os.path.exists(wmplayer_path):
                subprocess.Popen([wmplayer_path])
            else:
                subprocess.Popen("start mswindowsmusic:", shell=True)
            if song_name:
                return f"Windows Media Player open kar diya hai '{song_name}' ke liye."
            return "Windows Media Player open kar diya hai."


